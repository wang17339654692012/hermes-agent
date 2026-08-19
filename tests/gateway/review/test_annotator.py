"""Tests for annotator — Word 批注生成。"""

import sys
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))

from gateway.platforms.review.annotator import (
    generate_annotated_docx,
    _generate_sync,
    SEVERITY_LABELS,
    SEVERITY_LABEL_COLORS,
)
from gateway.platforms.review.models import Annotation
from docx.oxml.ns import qn


def _make_test_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段测试内容。")
    doc.add_paragraph("第二段测试内容，包含问题。")
    doc.add_paragraph("第三段正常内容。")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_test_annotations() -> list:
    return [
        Annotation(
            paragraph_index=2,
            severity="critical",
            issue_type="政治表述错误",
            description="测试描述",
            suggestion="修改建议",
            reference="参考来源",
        ),
        Annotation(
            paragraph_index=1,
            severity="suggestion",
            issue_type="措辞优化",
            description="可以更规范",
            suggestion="优化建议",
            reference="参考来源2",
        ),
    ]


class TestSeverityConstants:
    """常量测试."""

    def test_labels_defined(self):
        assert SEVERITY_LABELS["critical"] == "● 严重"
        assert SEVERITY_LABELS["important"] == "● 重要"
        assert SEVERITY_LABELS["suggestion"] == "● 建议"

    def test_label_colors_defined(self):
        assert SEVERITY_LABEL_COLORS["critical"] == "FF0000"
        assert SEVERITY_LABEL_COLORS["important"] == "BF8F00"
        assert SEVERITY_LABEL_COLORS["suggestion"] == "0070C0"

    def test_labels_no_emoji(self):
        """级别标签用 ●(U+25CF) + 着色，不用 emoji——emoji 在 WPS 缺字形显示方块。"""
        for label in SEVERITY_LABELS.values():
            assert "🔴" not in label
            assert "🟡" not in label
            assert "🔵" not in label
            assert "●" in label


class TestGenerateAnnotatedDocx:
    """generate_annotated_docx tests."""

    def test_generate_with_annotations(self):
        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = _generate_sync(file_bytes, "test.docx", annotations)
        assert isinstance(result, bytes)
        assert len(result) > 0
        # 应该比原始文件大（多了批注）
        assert len(result) > len(file_bytes)

    def test_generate_no_annotations(self):
        file_bytes = _make_test_docx()
        annotations = []

        result = _generate_sync(file_bytes, "test.docx", annotations)
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_generate_empty_document(self):
        from docx import Document

        doc = Document()
        buf = BytesIO()
        doc.save(buf)

        result = _generate_sync(buf.getvalue(), "test.docx", [])
        assert isinstance(result, bytes)

    def test_output_is_valid_docx(self):
        from docx import Document

        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = _generate_sync(file_bytes, "test.docx", annotations)

        # 验证可以重新打开
        doc = Document(BytesIO(result))
        assert len(doc.paragraphs) >= 3

    def test_no_summary_table(self):
        """汇总表已移除，文档中不应出现。"""
        from docx import Document

        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = _generate_sync(file_bytes, "test.docx", annotations)

        doc = Document(BytesIO(result))
        for para in doc.paragraphs:
            assert "审核汇总表" not in para.text, "汇总表不应出现"

    def test_comments_present(self):
        """批注应写入 comments.xml。"""
        import zipfile
        from lxml import etree

        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = _generate_sync(file_bytes, "test.docx", annotations)

        with zipfile.ZipFile(BytesIO(result), "r") as z:
            assert "word/comments.xml" in z.namelist()
            xml = z.read("word/comments.xml")
            root = etree.fromstring(xml)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            comments = root.findall(".//w:comment", ns)
            assert len(comments) == 2  # 2 annotations

    def test_comment_label_colored_no_emoji(self):
        """批注气泡首行（级别行）彩色圆点 + 着色，无 emoji。"""
        import zipfile

        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = _generate_sync(file_bytes, "test.docx", annotations)

        with zipfile.ZipFile(BytesIO(result), "r") as z:
            xml = z.read("word/comments.xml").decode("utf-8")

        # 级别行带 w:color（critical → FF0000），含圆点，无 emoji 字符
        assert 'w:val="FF0000"' in xml
        assert "● 严重" in xml
        for ch in ("🔴", "🟡", "🔵"):
            assert ch not in xml

    def test_no_paragraph_background(self):
        """正文段落不添加背景色/高亮（批注级别靠气泡标签区分）。"""
        import zipfile

        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = _generate_sync(file_bytes, "test.docx", annotations)

        with zipfile.ZipFile(BytesIO(result), "r") as z:
            xml = z.read("word/document.xml").decode("utf-8")

        assert "w:shd" not in xml
        assert "w:highlight" not in xml

    def test_duplicate_text_paragraphs_anchored_by_position(self):
        """相同文本的两个段落，批注按文档位置锚定而非文本匹配。

        回归测试：文本匹配会把第 2 段的批注错误落到第 1 段（或丢失）。
        """
        from docx import Document

        doc = Document()
        doc.add_paragraph("重复文本段落。")
        doc.add_paragraph("重复文本段落。")  # 与第 1 段完全相同
        doc.add_paragraph("独立段落。")
        buf = BytesIO()
        doc.save(buf)

        annotations = [
            Annotation(
                paragraph_index=2, severity="important", issue_type="T",
                description="d", suggestion="s", reference="r",
            ),
        ]

        result = _generate_sync(buf.getvalue(), "test.docx", annotations)

        out_doc = Document(BytesIO(result))
        second_para_el = out_doc.paragraphs[1]._element
        first_para_el = out_doc.paragraphs[0]._element
        assert second_para_el.find(qn("w:commentRangeStart")) is not None
        assert first_para_el.find(qn("w:commentRangeStart")) is None

    @pytest.mark.asyncio
    async def test_async_wrapper(self):
        file_bytes = _make_test_docx()
        annotations = _make_test_annotations()

        result = await generate_annotated_docx(
            file_bytes, "test.docx", annotations
        )
        assert isinstance(result, bytes)
        assert len(result) > 0