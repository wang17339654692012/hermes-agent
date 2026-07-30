"""Tests for parser — 文档解析。"""

import sys
from io import BytesIO
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))

from gateway.platforms.review.parser import parse_document, Paragraph, _parse_docx, _parse_pdf


class TestParseDocument:
    """parse_document tests."""

    def test_docx_parsing(self):
        from docx import Document

        doc = Document()
        doc.add_heading("通知标题", level=1)
        doc.add_paragraph("第一段内容。")
        doc.add_paragraph("第二段内容。")
        doc.add_paragraph("")  # 空段落

        buf = BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        paragraphs = parse_document(file_bytes, "test.docx")
        assert len(paragraphs) == 3  # 空段落被跳过
        assert paragraphs[0].text == "通知标题"
        assert paragraphs[0].style == "Heading 1"
        assert paragraphs[1].text == "第一段内容。"
        assert paragraphs[2].text == "第二段内容。"

    def test_docx_empty(self):
        from docx import Document

        doc = Document()
        buf = BytesIO()
        doc.save(buf)

        paragraphs = parse_document(buf.getvalue(), "test.docx")
        assert len(paragraphs) == 0

    def test_pdf_parsing(self):
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "第一段内容。\n\n第二段内容。", fontsize=12)
        buf = BytesIO()
        doc.save(buf)
        doc.close()

        paragraphs = parse_document(buf.getvalue(), "test.pdf")
        assert len(paragraphs) >= 1

    def test_unsupported_format(self):
        with pytest.raises(ValueError, match="Unsupported format"):
            parse_document(b"dummy", "test.txt")

    def test_paragraph_index(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("段落1")
        doc.add_paragraph("段落2")
        doc.add_paragraph("段落3")

        buf = BytesIO()
        doc.save(buf)

        paragraphs = parse_document(buf.getvalue(), "test.docx")
        assert paragraphs[0].index == 1
        assert paragraphs[1].index == 2
        assert paragraphs[2].index == 3

    def test_runs_preserved(self):
        from docx import Document

        doc = Document()
        para = doc.add_paragraph()
        para.add_run("加粗文字").bold = True
        para.add_run("普通文字")

        buf = BytesIO()
        doc.save(buf)

        paragraphs = parse_document(buf.getvalue(), "test.docx")
        assert paragraphs[0].runs is not None
        assert len(paragraphs[0].runs) == 2

    def test_paragraph_style(self):
        from docx import Document

        doc = Document()
        doc.add_heading("标题", level=1)
        doc.add_paragraph("正文")

        buf = BytesIO()
        doc.save(buf)

        paragraphs = parse_document(buf.getvalue(), "test.docx")
        assert paragraphs[0].style == "Heading 1"
        assert paragraphs[1].style == "Normal"