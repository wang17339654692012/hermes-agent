"""生成测试样例公文 — 含故意错误的《关于开展安全生产专项检查的通知》。

在 Hermes 镜像内执行（review extra 已烘焙 python-docx/pymupdf）：
  docker run --rm -v "$(pwd)/deploy/review/samples:/out" <img> \
      /opt/hermes/.venv/bin/python /opt/hermes/deploy/review/tools/generate_sample_docx.py --out /out

产出 sample_notice.docx 与 sample_notice.pdf（同主题内容）。
故意植入 6 类错误，对应 document-review 技能的三级审核标准：
  1. 固定政治表述错字（critical）：「习新平」
  2. 讲话引用场合错误（critical）：「建党 90 周年大会」实为 100 周年
  3. 文件名残缺（critical）：《中共中央关于党的百年奋斗和历史经验的决议》漏「重大成就」
  4. 提法偏差（important）：「二个确立」
  5. 格式要素缺失（important）：落款缺成文日期、主送机关缺全角冒号
  6. 口语化用词（suggestion）：「搞定」「立马」
"""

import argparse
import sys
from pathlib import Path


def build_docx(out_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("关于开展安全生产专项检查的通知", level=1)

    # 主送机关：故意用半角冒号（应为全角）
    doc.add_paragraph("各科室:")

    doc.add_heading("一、检查目的", level=2)
    doc.add_paragraph(
        "为深入贯彻习新平新时代中国特色社会主义思想，"
        "牢固树立安全发展理念，坚决防范和遏制各类安全事故发生，"
        "经研究决定，在全单位范围内开展安全生产专项检查。"
    )

    doc.add_heading("二、检查内容", level=2)
    doc.add_paragraph(
        "（一）落实安全生产责任制情况。各级各部门要深刻领悟二个确立的决定性意义，"
        "增强四个意识、坚定四个自信、做到两个维护，把安全生产责任扛在肩上。"
    )
    doc.add_paragraph(
        "（二）重点部位隐患排查情况。对办公区域、仓库、配电室、食堂等部位逐一排查，"
        "建立隐患台账，明确整改责任人和整改时限。"
    )

    doc.add_heading("三、工作要求", level=2)
    doc.add_paragraph(
        "（一）提高政治站位。习近平总书记在庆祝中国共产党成立90周年大会上的"
        "重要讲话为做好新时代各项工作提供了根本遵循，全体人员要认真学习领会。"
    )
    doc.add_paragraph(
        "（二）严格对照《中共中央关于党的百年奋斗和历史经验的决议》要求，"
        "查摆问题、补齐短板。"
    )
    doc.add_paragraph(
        "（三）各科室要真抓实干，把安全工作搞定，检查中发现的问题要立马整改，"
        "确保专项检查取得实效。"
    )

    doc.add_paragraph("特此通知。")

    # 落款：故意缺成文日期
    doc.add_paragraph("XX市XX局办公室")

    doc.save(out_path)


def build_pdf(out_path: Path, text_lines: list[str]) -> None:
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_font(fontname="china-s")
    y = 72
    for line in text_lines:
        page.insert_text((72, y), line, fontname="china-s", fontsize=12)
        y += 20
    pdf.save(out_path)
    pdf.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", required=True, help="输出目录")
    args = parser.parse_args()

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    docx_path = out_dir / "sample_notice.docx"
    build_docx(docx_path)

    # PDF 与 docx 同主题（PyMuPDF 不支持直接打开 docx，用内置 CJK 字体写入文本）
    text_lines = [
        "关于开展安全生产专项检查的通知",
        "各科室:",
        "为深入贯彻习新平新时代中国特色社会主义思想，牢固树立安全发展理念，",
        "坚决防范和遏制各类安全事故发生，经研究决定开展安全生产专项检查。",
        "各级各部门要深刻领悟二个确立的决定性意义，增强四个意识、坚定四个自信。",
        "习近平总书记在庆祝中国共产党成立90周年大会上的重要讲话为做好新时代",
        "各项工作提供了根本遵循，全体人员要认真学习领会。",
        "各科室要真抓实干，把安全工作搞定，检查中发现的问题要立马整改。",
        "特此通知。",
        "XX市XX局办公室",
    ]
    pdf_path = out_dir / "sample_notice.pdf"
    build_pdf(pdf_path, text_lines)

    print(f"[generate-sample] {docx_path} ({docx_path.stat().st_size} bytes)")
    print(f"[generate-sample] {pdf_path} ({pdf_path.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
