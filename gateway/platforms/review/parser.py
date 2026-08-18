"""文档解析 — .docx / .pdf → 结构化段落列表。

docx：解析 Word 元数据（自动编号 numbering.xml、样式、outlineLvl），
     调用 docmodel 渲染编号并做结构定位（角色/区域）。
pdf：无 Word 元数据，用启发式填充同一模型（审核能力自然降级，
     不另写一套逻辑）。

段落 index 与 python-docx doc.paragraphs 的枚举序号一致（含空段占位），
annotator 据此锚定批注位置。
"""

from io import BytesIO
from pathlib import Path
import logging
from typing import List, Optional

from . import docmodel
from .docmodel import NumberingRenderer, enrich
from .models import Paragraph

logger = logging.getLogger(__name__)


def parse_document(file_bytes: bytes, filename: str) -> List[Paragraph]:
    """解析文档为段落列表。

    Args:
        file_bytes: 文档二进制内容
        filename: 原始文件名（用于判断扩展名）

    Returns:
        Paragraph 列表（已渲染自动编号、已填充 role/region）

    Raises:
        ValueError: 不支持的文档格式
    """
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        return _parse_docx(file_bytes)
    elif ext == ".pdf":
        return _parse_pdf(file_bytes)
    else:
        raise ValueError(f"Unsupported format: {ext}. Use .docx or .pdf")


def _parse_docx(file_bytes: bytes) -> List[Paragraph]:
    """解析 .docx：自动编号渲染 + 样式 + 结构定位。"""
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn

    doc = DocxDocument(BytesIO(file_bytes))

    # 读取 numbering.xml（Word 列表编号定义）
    numbering_xml: Optional[bytes] = None
    try:
        for part in doc.part.package.iter_parts():
            if str(part.partname) == "/word/numbering.xml":
                numbering_xml = part.blob
                break
    except Exception as e:
        logger.warning("读取 numbering.xml 失败: %s", e)

    renderer = NumberingRenderer(numbering_xml)

    paragraphs: List[Paragraph] = []
    for i, para in enumerate(doc.paragraphs, start=1):
        # 渲染自动编号：numPr → 编号前缀
        number = ""
        try:
            p_pr = para._p.find(qn("w:pPr"))
            if p_pr is not None:
                num_pr = p_pr.find(qn("w:numPr"))
                if num_pr is not None:
                    ilvl_el = num_pr.find(qn("w:ilvl"))
                    num_id_el = num_pr.find(qn("w:numId"))
                    ilvl = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
                    num_id = int(num_id_el.get(qn("w:val"))) if num_id_el is not None else None
                    number = renderer.render(num_id, ilvl)
        except (TypeError, ValueError):
            number = ""

        text = para.text.strip()
        if not text and not number:
            continue  # 跳过空段落（空编号段落保留渲染编号）

        # outlineLvl（大纲级别，可用于层级判定）
        outline_level = None
        try:
            if p_pr is not None:
                ol_el = p_pr.find(qn("w:outlineLvl"))
                if ol_el is not None:
                    outline_level = int(ol_el.get(qn("w:val")))
        except (TypeError, ValueError):
            outline_level = None

        paragraphs.append(Paragraph(
            index=i,
            text=text,  # 原始文本；编号前缀由 enrich() 统一拼接一次
            style=para.style.name if para.style else "Normal",
            runs=list(para.runs),
            number=number,
            outline_level=outline_level,
        ))

    return enrich(paragraphs, renderer)


def _parse_pdf(file_bytes: bytes) -> List[Paragraph]:
    """解析 PDF，按段落分割；启发式结构定位（无 Word 元数据）。"""
    import fitz  # pymupdf

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    paragraphs: List[Paragraph] = []
    para_index = 0

    for page in doc:
        text = page.get_text("text")
        for block in text.split("\n\n"):
            block = block.strip()
            if block:
                para_index += 1
                paragraphs.append(Paragraph(
                    index=para_index,
                    text=block,
                    style="Normal",
                ))

    return enrich(paragraphs, None)
